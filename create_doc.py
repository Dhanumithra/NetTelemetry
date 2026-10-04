import sys
from docx import Document

def create_doc(filename, title, objective, what_done, result):
    doc = Document()
    doc.add_heading(title, 0)
    
    doc.add_heading('Objective', level=1)
    doc.add_paragraph(objective)
    
    doc.add_heading('What is Done & Why', level=1)
    doc.add_paragraph(what_done)
    
    doc.add_heading('Result', level=1)
    doc.add_paragraph(result)
    
    doc.save(filename)
    print(f"Created {filename}")

if __name__ == "__main__":
    import json
    with open("doc_args.json", "r") as f:
        args = json.load(f)
    create_doc(args["filename"], args["title"], args["objective"], args["what_done"], args["result"])
